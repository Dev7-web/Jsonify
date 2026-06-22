import asyncio
from pathlib import Path
import sys
from typing import Any

from openpyxl import Workbook


sys.path.append(str(Path(__file__).resolve().parents[1]))

from app.models import Document, DocumentSchema
from app.parsing.excel_loader import load_sheets
from app.parsing.form_detector import detect_key_value_sections
from app.parsing.header_detector import detect_table_headers
from app.pipeline import detect as detect_pipeline
from app.pipeline.fingerprint import build_deterministic_layout_fingerprint
from app.pipeline.fingerprint import build_header_structure_fingerprint


BACKEND_ROOT = Path(__file__).resolve().parents[1]
UPLOADS_DIR = BACKEND_ROOT / "uploads"
WORKBOOK_NAME = "step14-short-circuit.xlsx"
STORED_PATH = f"uploads/{WORKBOOK_NAME}"


APPROVED_HEADER_STRUCTURE = {
    "sheets": [
        {
            "name": "Original Lease - Lease",
            "sections": [
                {
                    "type": "key_value",
                    "title": "Amendment Abstract",
                    "fields": ["Lease ID", "DBA", "Type", "Lease", "Status"],
                },
                {
                    "type": "table",
                    "title": "Space",
                    "headers": [{"name": "Unit"}, {"name": "Building"}, {"name": "Area"}],
                },
                {
                    "type": "table",
                    "title": "Charge Schedules",
                    "headers": [
                        {"name": "Charge Code"},
                        {"name": "Charge Desc"},
                        {"name": "Date From"},
                        {"name": "Date To"},
                        {"name": "Monthly Amt"},
                        {"name": "Annual Amt"},
                    ],
                },
            ],
        }
    ]
}


class FakeUpdateResult:
    def __init__(self, matched_count: int) -> None:
        self.matched_count = matched_count


class FakeCursor:
    def __init__(self, records: list[dict[str, Any]]) -> None:
        self.records = records
        self.index = 0

    def __aiter__(self):
        return self

    async def __anext__(self) -> dict[str, Any]:
        if self.index >= len(self.records):
            raise StopAsyncIteration

        record = self.records[self.index]
        self.index += 1
        return record


class FakeCollection:
    def __init__(self, records: list[dict[str, Any]]) -> None:
        self.records = {record["_id"]: record for record in records}

    async def find_one(self, query: dict[str, Any]) -> dict[str, Any] | None:
        record_id = query.get("_id")
        if record_id is None:
            return None

        record = self.records.get(record_id)
        if record is None:
            return None

        expected_status = query.get("status")
        if isinstance(expected_status, dict) and "$ne" in expected_status:
            if record.get("status") == expected_status["$ne"]:
                return None
        elif expected_status is not None and record.get("status") != expected_status:
            return None

        return record

    async def update_one(
        self,
        query: dict[str, Any],
        update: dict[str, Any],
    ) -> FakeUpdateResult:
        record = await self.find_one(query)
        if record is None:
            return FakeUpdateResult(0)

        for key, value in update.get("$set", {}).items():
            record[key] = value

        for key in update.get("$unset", {}):
            record.pop(key, None)

        return FakeUpdateResult(1)

    def find(self, query: dict[str, Any]) -> FakeCursor:
        records = []
        for record in self.records.values():
            if all(record.get(key) == value for key, value in query.items()):
                records.append(record)

        return FakeCursor(records)


class FakeDb:
    def __init__(self, document: Document, schemas: list[DocumentSchema]) -> None:
        self.documents = FakeCollection([document.model_dump(by_alias=True)])
        self.schemas = FakeCollection(
            [schema.model_dump(by_alias=True) for schema in schemas]
        )


def create_workbook(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Original Lease - Lease"

    worksheet["A1"] = "Space"
    worksheet["A2"] = "Unit"
    worksheet["B2"] = "Building"
    worksheet["C2"] = "Area"
    worksheet["A3"] = "101"
    worksheet["B3"] = "Main"
    worksheet["C3"] = 1200

    worksheet["A5"] = "Charge Schedules"
    worksheet["A6"] = "Charge Code"
    worksheet["B6"] = "Charge Desc"
    worksheet["C6"] = "Date From"
    worksheet["D6"] = "Date To"
    worksheet["E6"] = "Monthly Amt"
    worksheet["F6"] = "Annual Amt"
    worksheet["A7"] = "BASE"
    worksheet["B7"] = "Base Rent"
    worksheet["C7"] = "2026-01-01"
    worksheet["D7"] = "2026-12-31"
    worksheet["E7"] = 1000
    worksheet["F7"] = 12000

    workbook.save(path)


def create_deterministic_fingerprint(path: Path) -> str:
    deterministic_headers = {
        sheet_name: detect_pipeline._deduplicate_headers(detect_table_headers(grid))
        for sheet_name, grid in load_sheets(path).items()
    }
    deterministic_key_value_sections = detect_key_value_sections(path)

    return build_deterministic_layout_fingerprint(
        deterministic_headers,
        deterministic_key_value_sections,
    ).fingerprint


async def fail_if_llm_is_called(*_: object, **__: object):
    raise AssertionError("LLM should not be called when schema fingerprint matches.")


async def check_schema_match_skips_llm() -> None:
    workbook_path = UPLOADS_DIR / WORKBOOK_NAME
    create_workbook(workbook_path)
    deterministic_fingerprint = create_deterministic_fingerprint(workbook_path)

    document = Document(
        filename=WORKBOOK_NAME,
        file_type="xlsx",
        stored_path=STORED_PATH,
    )
    schema = DocumentSchema(
        name="Dollar Tree approved layout",
        file_type="xlsx",
        fingerprint=build_header_structure_fingerprint(APPROVED_HEADER_STRUCTURE),
        deterministic_fingerprint=deterministic_fingerprint,
        version=1,
        status="active",
        header_structure=APPROVED_HEADER_STRUCTURE,
    )
    fake_db = FakeDb(document, [schema])

    original_llm = detect_pipeline.adetect_header_structure_from_cell_map
    detect_pipeline.adetect_header_structure_from_cell_map = fail_if_llm_is_called

    try:
        result = await detect_pipeline.detect_document_headers(document.id, db=fake_db)
    finally:
        detect_pipeline.adetect_header_structure_from_cell_map = original_llm
        workbook_path.unlink(missing_ok=True)

    stored_document = fake_db.documents.records[document.id]

    assert result["source"] == "schema"
    assert result["status"] == "approved"
    assert result["matched_schema_id"] == schema.id
    assert stored_document["status"] == "approved"
    assert stored_document["matched_schema_id"] == schema.id
    assert stored_document["detected_headers"]["source"] == "schema"

    print("Schema short-circuit branch passed.")
    print(f"Matched schema: {schema.id}")
    print(f"Source: {result['source']}")


async def check_no_match_uses_llm() -> None:
    workbook_path = UPLOADS_DIR / WORKBOOK_NAME
    create_workbook(workbook_path)

    document = Document(
        filename=WORKBOOK_NAME,
        file_type="xlsx",
        stored_path=STORED_PATH,
    )
    fake_db = FakeDb(document, [])
    llm_calls = 0

    async def fake_llm(*_: object, **__: object):
        nonlocal llm_calls
        llm_calls += 1
        return APPROVED_HEADER_STRUCTURE

    original_llm = detect_pipeline.adetect_header_structure_from_cell_map
    detect_pipeline.adetect_header_structure_from_cell_map = fake_llm

    try:
        result = await detect_pipeline.detect_document_headers(document.id, db=fake_db)
    finally:
        detect_pipeline.adetect_header_structure_from_cell_map = original_llm
        workbook_path.unlink(missing_ok=True)

    stored_document = fake_db.documents.records[document.id]

    assert llm_calls == 1
    assert result["source"] == "llm"
    assert result["status"] == "needs_review"
    assert stored_document["status"] == "needs_review"
    assert stored_document["detected_headers"]["source"] == "llm"

    print("LLM fallback branch passed.")
    print(f"Source: {result['source']}")


async def main() -> None:
    await check_schema_match_skips_llm()
    await check_no_match_uses_llm()
    print("Step 14 schema matching check passed.")


if __name__ == "__main__":
    asyncio.run(main())
