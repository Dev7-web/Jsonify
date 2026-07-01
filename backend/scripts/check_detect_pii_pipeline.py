from __future__ import annotations

import asyncio
import os
from pathlib import Path
import sys
from typing import Any

from openpyxl import Workbook


sys.path.append(str(Path(__file__).resolve().parents[1]))

import app.llm.client as llm_client
import app.pipeline.detect as detect_pipeline
from app.models import Document
from app.pii.adapt import contains_placeholder
from app.pipeline.detect import detect_document_headers


BACKEND_ROOT = Path(__file__).resolve().parents[1]
UPLOADS_DIR = BACKEND_ROOT / "uploads"
WORKBOOK_NAME = "pii-detect-smoke.xlsx"
STORED_PATH = f"uploads/{WORKBOOK_NAME}"


class FakeUpdateResult:
    def __init__(self, matched_count: int) -> None:
        self.matched_count = matched_count


class FakeCollection:
    def __init__(self, records: list[dict[str, Any]]) -> None:
        self.records = {record["_id"]: record for record in records}

    async def find_one(self, query: dict[str, Any]) -> dict[str, Any] | None:
        record_id = query.get("_id")
        if record_id is None:
            return None

        return self.records.get(record_id)

    async def update_one(
        self,
        query: dict[str, Any],
        update: dict[str, Any],
    ) -> FakeUpdateResult:
        record = await self.find_one(query)
        if record is None:
            return FakeUpdateResult(0)

        status_filter = query.get("status")
        if isinstance(status_filter, dict) and status_filter.get("$ne") == record.get("status"):
            return FakeUpdateResult(0)

        for key, value in update.get("$set", {}).items():
            record[key] = value

        for key in update.get("$unset", {}):
            record.pop(key, None)

        return FakeUpdateResult(1)


class FakeDb:
    def __init__(self, document: Document) -> None:
        self.documents = FakeCollection([document.model_dump(by_alias=True)])
        self.schemas = FakeCollection([])


def create_workbook(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Sheet1"
    worksheet.append(["Contact", "Contact"])
    worksheet.append(["Name", "Neha Rao"])
    worksheet.append(["Email", "neha@acme.in"])
    worksheet.append([])
    worksheet.append(["Charge Schedule", "Charge Schedule"])
    worksheet.append(["Code", "Amount"])
    worksheet.append(["BASE", 100])
    workbook.save(path)


async def main() -> None:
    await check_post_parse_restore_before_hints()

    workbook_path = UPLOADS_DIR / WORKBOOK_NAME
    create_workbook(workbook_path)

    document = Document(
        filename=WORKBOOK_NAME,
        file_type="xlsx",
        stored_path=STORED_PATH,
    )
    fake_db = FakeDb(document)
    captured_prompts: list[str] = []

    async def fake_gemini(_: str, user: str, **__: object) -> str:
        captured_prompts.append(user)
        return (
            '{"sheets":[{"name":"Sheet1","sections":['
            '{"type":"key_value","title":"Contact","fields":["Name","Email"]},'
            '{"type":"table","title":"Charge Schedule","headers":['
            '{"name":"Code"},{"name":"Amount"}]}'
            "]}]}"
        )

    async def fake_pii_redact(payload: Any) -> tuple[str, Any]:
        redacted = {
            key: (
                "<PERSON_1234>"
                if value == "Neha Rao"
                else "<EMAIL_1234>"
                if value == "neha@acme.in"
                else value
            )
            for key, value in payload.items()
        }
        return "pii-doc-smoke", redacted

    original_gemini = llm_client.gemini_call
    original_pii_redact = llm_client.pii_redact
    original_pii_enabled = os.environ.get("PII_ENABLED")
    os.environ["PII_ENABLED"] = "true"
    llm_client.gemini_call = fake_gemini
    llm_client.pii_redact = fake_pii_redact

    try:
        result = await detect_document_headers(document.id, db=fake_db)
    finally:
        llm_client.gemini_call = original_gemini
        llm_client.pii_redact = original_pii_redact
        if original_pii_enabled is None:
            os.environ.pop("PII_ENABLED", None)
        else:
            os.environ["PII_ENABLED"] = original_pii_enabled
        workbook_path.unlink(missing_ok=True)

    stored_document = fake_db.documents.records[document.id]
    assert result["source"] == "llm"
    assert stored_document["status"] == "needs_review"
    assert stored_document["pii_document_id"]
    assert captured_prompts
    assert "Neha Rao" not in captured_prompts[0]
    assert "neha@acme.in" not in captured_prompts[0]
    assert contains_placeholder(captured_prompts[0])

    print("Detect PII pipeline check passed.")
    print("PII document id stored:", bool(stored_document["pii_document_id"]))


async def check_post_parse_restore_before_hints() -> None:
    async def fake_pii_restore(document_id: str, payload: Any) -> Any:
        assert document_id == "pii-doc-restore"
        assert payload["sheets"][0]["sections"][0]["title"] == (
            "<ORGANIZATION_bf09> Information"
        )
        return {
            "sheets": [
                {
                    "name": "Renewal - 1st Renewal",
                    "sections": [
                        {
                            "type": "key_value",
                            "title": "Lease Information",
                            "fields": ["DBA", "Lease"],
                        }
                    ],
                }
            ]
        }

    original_pii_restore = detect_pipeline.pii_restore
    detect_pipeline.pii_restore = fake_pii_restore

    try:
        restored = await detect_pipeline._restore_header_structure_placeholders(
            {
                "sheets": [
                    {
                        "name": "Renewal - 1st Renewal",
                        "sections": [
                            {
                                "type": "key_value",
                                "title": "<ORGANIZATION_bf09> Information",
                                "fields": ["DBA", "Lease"],
                            }
                        ],
                    }
                ]
            },
            "pii-doc-restore",
        )
    finally:
        detect_pipeline.pii_restore = original_pii_restore

    merged = detect_pipeline.apply_key_value_section_hints(
        restored,
        {
            "Renewal - 1st Renewal": [
                {
                    "title": "Lease Information",
                    "title_coordinate": "V4",
                    "row_index": 4,
                    "coordinate": "V4",
                    "fields": [
                        {"name": "DBA", "coordinate": "V5"},
                        {"name": "Lease", "coordinate": "V6"},
                    ],
                }
            ]
        },
    )

    sections = merged["sheets"][0]["sections"]
    assert len(sections) == 1
    assert sections[0]["title"] == "Lease Information"
    detect_pipeline._raise_if_header_structure_has_placeholders(merged)

    fallback_merged = detect_pipeline.apply_key_value_section_hints(
        {
            "sheets": [
                {
                    "name": "Renewal - 1st Renewal",
                    "sections": [
                        {
                            "type": "key_value",
                            "title": "<ORGANIZATION_bf09> Information",
                            "fields": ["DBA", "Lease"],
                        }
                    ],
                }
            ]
        },
        {
            "Renewal - 1st Renewal": [
                {
                    "title": "Lease Information",
                    "title_coordinate": "V4",
                    "row_index": 4,
                    "coordinate": "V4",
                    "fields": [
                        {"name": "DBA", "coordinate": "V5"},
                        {"name": "Lease", "coordinate": "V6"},
                    ],
                }
            ]
        },
    )
    fallback_sections = fallback_merged["sheets"][0]["sections"]
    assert len(fallback_sections) == 1
    assert fallback_sections[0]["title"] == "Lease Information"
    detect_pipeline._raise_if_header_structure_has_placeholders(fallback_merged)


if __name__ == "__main__":
    asyncio.run(main())
