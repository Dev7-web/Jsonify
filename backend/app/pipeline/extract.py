from __future__ import annotations

from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path
import re
from typing import Any, TypedDict

from motor.motor_asyncio import AsyncIOMotorDatabase
from openpyxl.utils import get_column_letter

from app.config import BACKEND_ROOT
from app.db import get_db
from app.models import Document, DocumentSchema
from app.parsing.excel_loader import Grid, MergedRanges, load_sheets_with_merged_ranges
from app.parsing.matrix_detector import DetectedMatrix, detect_matrices


FieldLocators = dict[str, Any]
ExtractedJson = dict[str, Any]

_RESOLVED_BACKEND_ROOT = BACKEND_ROOT.resolve()
_WHITESPACE_PATTERN = re.compile(r"\s+")
_PLACEHOLDER_TOKEN_PATTERN = re.compile(r"\{\{[^{}]*\}\}")
_ALNUM_PATTERN = re.compile(r"[^\W_]")
_SYNTHETIC_SECTION_TITLE_PATTERN = re.compile(
    r".+\s+section\s+\d+\s*$",
    re.IGNORECASE,
)
_LOCATOR_VERSION = 7
_LOCATOR_TYPE = "excel"
_MAX_SECTION_SEARCH_ROWS = 80
_REBO_COMMENT_OUTPUT_TITLE = "REBO Comment"
_REBO_COMMENT_NORMALIZED_TITLES = {"rebo comment", "rebo comments"}


class ExtractionResponse(TypedDict):
    id: str
    status: str
    schema_id: str
    output_json: ExtractedJson
    field_locators: FieldLocators
    locators_created: bool


class StoredJsonResponse(TypedDict):
    id: str
    status: str
    output_json: ExtractedJson


class ExtractError(RuntimeError):
    pass


class ExtractDocumentNotFoundError(ExtractError):
    pass


class ExtractSchemaNotFoundError(ExtractError):
    pass


class ExtractUnsupportedDocumentTypeError(ExtractError):
    pass


class ExtractStoredDocumentFileError(ExtractError):
    pass


class ExtractInvalidDocumentFileError(ExtractError):
    pass


class ExtractMissingSchemaLinkError(ExtractError):
    pass


class ExtractLocatorError(ExtractError):
    pass


class ExtractOutputNotFoundError(ExtractError):
    pass


async def extract_document_json(
    document_id: str,
    *,
    db: AsyncIOMotorDatabase | None = None,
) -> ExtractionResponse:
    database = db if db is not None else get_db()

    document_record = await database.documents.find_one({"_id": document_id})
    if document_record is None:
        raise ExtractDocumentNotFoundError(f"Document not found: {document_id}")

    document = Document.model_validate(document_record)
    if document.file_type != "xlsx":
        raise ExtractUnsupportedDocumentTypeError("Only .xlsx extraction is supported.")

    if not document.matched_schema_id:
        raise ExtractMissingSchemaLinkError(
            "Document must be approved and linked to a schema before extraction."
        )

    schema_record = await database.schemas.find_one({"_id": document.matched_schema_id})
    if schema_record is None:
        raise ExtractSchemaNotFoundError(
            f"Schema not found: {document.matched_schema_id}"
        )
    schema = DocumentSchema.model_validate(schema_record)

    workbook_path = _resolve_stored_path(document.stored_path)
    sheets, merged_ranges_by_sheet = _load_workbook_sheets(workbook_path)

    locators_created = False
    field_locators = document.extraction_locators
    if not _is_supported_document_field_locators(field_locators, schema.id):
        field_locators = build_excel_field_locators(
            schema.header_structure,
            sheets,
            schema_id=schema.id,
            merged_ranges_by_sheet=merged_ranges_by_sheet,
        )
        locators_created = True

    output_json = extract_excel_with_locators(sheets, field_locators)
    update_result = await database.documents.update_one(
        {"_id": document.id},
        {
            "$set": {
                "extraction_locators": field_locators,
                "output_json": output_json,
                "status": "extracted",
            },
            "$unset": {"failure_reason": ""},
        },
    )
    if update_result.matched_count == 0:
        raise ExtractDocumentNotFoundError(f"Document not found: {document_id}")

    return {
        "id": document.id,
        "status": "extracted",
        "schema_id": schema.id,
        "output_json": output_json,
        "field_locators": field_locators,
        "locators_created": locators_created,
    }


async def get_document_output_json(
    document_id: str,
    *,
    db: AsyncIOMotorDatabase | None = None,
) -> StoredJsonResponse:
    database = db if db is not None else get_db()

    document_record = await database.documents.find_one({"_id": document_id})
    if document_record is None:
        raise ExtractDocumentNotFoundError(f"Document not found: {document_id}")

    document = Document.model_validate(document_record)
    if document.output_json is None:
        raise ExtractOutputNotFoundError(
            f"Document has no extracted JSON yet: {document_id}"
        )

    return {
        "id": document.id,
        "status": document.status,
        "output_json": document.output_json,
    }


def build_excel_field_locators(
    header_structure: dict[str, Any],
    sheets: dict[str, Grid],
    *,
    schema_id: str | None = None,
    merged_ranges_by_sheet: MergedRanges | None = None,
) -> FieldLocators:
    sheet_locators: list[dict[str, Any]] = []

    for sheet in header_structure.get("sheets", []):
        requested_sheet_name = _require_string(sheet.get("name"), "sheet.name")
        sheet_name, grid = _find_sheet_grid(sheets, requested_sheet_name)
        raw_sections = sheet.get("sections", [])
        if not isinstance(raw_sections, list):
            raise ExtractLocatorError(
                f'Sheet "{requested_sheet_name}" sections must be a list.'
            )

        sections = _workbook_supported_sections(grid, raw_sections)
        section_title_rows = _collect_section_title_rows(grid, sections)
        located_sections: list[dict[str, Any]] = []

        for section_index, section in enumerate(sections):
            if not isinstance(section, dict):
                raise ExtractLocatorError(
                    f'Section {section_index + 1} on sheet "{sheet_name}" must be an object.'
                )

            section_type = section.get("type")
            if section_type == "key_value":
                located_sections.append(
                    _locate_key_value_section(
                        grid,
                        section,
                        section_index=section_index,
                        section_title_rows=section_title_rows,
                    )
                )
                continue

            if section_type == "table":
                located_sections.append(
                    _locate_table_section(
                        grid,
                        section,
                        section_index=section_index,
                        section_title_rows=section_title_rows,
                    )
                )
                continue

            if section_type == "matrix":
                located_sections.append(
                    _locate_matrix_section(
                        grid,
                        section,
                        section_index=section_index,
                        section_title_rows=section_title_rows,
                        merged_ranges=(merged_ranges_by_sheet or {}).get(
                            sheet_name,
                            [],
                        ),
                    )
                )
                continue

            raise ExtractLocatorError(
                f'Section {section_index + 1} on sheet "{sheet_name}" has unsupported type.'
            )

        sheet_locators.append(
            {
                "name": sheet_name,
                "source_name": requested_sheet_name,
                "sections": located_sections,
            }
        )

    locators: FieldLocators = {
        "version": _LOCATOR_VERSION,
        "type": _LOCATOR_TYPE,
        "sheets": sheet_locators,
    }
    if schema_id is not None:
        locators["schema_id"] = schema_id

    return locators


def extract_excel_with_locators(
    sheets: dict[str, Grid],
    field_locators: FieldLocators,
) -> ExtractedJson:
    if not _is_supported_field_locators(field_locators):
        raise ExtractLocatorError("Unsupported or missing Excel field locators.")

    output: ExtractedJson = {}

    for sheet_locator in field_locators["sheets"]:
        sheet_name = _require_string(sheet_locator.get("name"), "locator.sheet.name")
        _, grid = _find_sheet_grid(sheets, sheet_name)
        sheet_output: dict[str, Any] = {}

        for section_locator in sheet_locator.get("sections", []):
            section_title = _section_output_title(section_locator)
            output_key = _unique_key(sheet_output, section_title)

            kind = section_locator.get("type")
            if kind == "key_value":
                if _is_rebo_comment_section(section_locator):
                    sheet_output[output_key] = _extract_rebo_comment_section(
                        grid,
                        section_locator,
                    )
                else:
                    sheet_output[output_key] = _extract_key_value_section(
                        grid,
                        section_locator,
                    )
            elif kind == "table":
                sheet_output[output_key] = _extract_table_section(
                    grid,
                    section_locator,
                )
            elif kind == "matrix":
                sheet_output[output_key] = _extract_matrix_section(
                    grid,
                    section_locator,
                )
            else:
                raise ExtractLocatorError(
                    f'Unsupported section locator type on sheet "{sheet_name}".'
                )

        output[sheet_name] = sheet_output

    return output


def _locate_key_value_section(
    grid: Grid,
    section: dict[str, Any],
    *,
    section_index: int,
    section_title_rows: dict[int, int],
) -> dict[str, Any]:
    title = _optional_string(section.get("title"))
    start_row, end_row = _section_search_bounds(
        section_index,
        section_title_rows,
        row_count=len(grid),
    )
    title_cell = (
        _find_section_title_cell(grid, title, start_row=start_row, end_row=end_row)
        if title
        else None
    )

    fields = section.get("fields", [])
    if not isinstance(fields, list):
        section_name = title or section_index + 1
        raise ExtractLocatorError(
            f'Key/value section "{section_name}" fields must be a list.'
        )

    field_names = [_require_string(field, "key_value field") for field in fields]
    section_field_labels = {_normalize_label(field_name) for field_name in field_names}
    field_locators = []
    for field_name in field_names:
        label_cell = _find_label_cell(
            grid,
            field_name,
            start_row=start_row,
            end_row=end_row,
        )

        if label_cell is None:
            field_locators.append(
                {
                    "name": field_name,
                    "status": "missing_label",
                    "label": field_name,
                }
            )
            continue

        value_cell = _find_value_cell_to_right(
            grid,
            row_index=label_cell["row_index"],
            column_index=label_cell["column_index"],
            label=field_name,
            stop_labels=section_field_labels - {_normalize_label(field_name)},
        )
        field_locators.append(
            {
                "name": field_name,
                "status": "found" if value_cell is not None else "missing_value",
                "label_coordinate": label_cell["coordinate"],
                "label_row_index": label_cell["row_index"],
                "label_column_index": label_cell["column_index"],
                **(
                    {
                        "value_coordinate": value_cell["coordinate"],
                        "value_row_index": value_cell["row_index"],
                        "value_column_index": value_cell["column_index"],
                    }
                    if value_cell is not None
                    else {}
                ),
            }
        )

    return {
        "type": "key_value",
        "title": title,
        "output_title": _key_value_section_output_title(
            grid,
            title,
            field_locators,
        ),
        "section_index": section_index,
        "search_start_row": start_row,
        "search_end_row": end_row,
        **(
            {
                "title_coordinate": title_cell["coordinate"],
                "title_row_index": title_cell["row_index"],
                "title_column_index": title_cell["column_index"],
            }
            if title_cell is not None
            else {}
        ),
        "fields": field_locators,
    }


def _workbook_supported_sections(
    grid: Grid,
    sections: list[Any],
) -> list[Any]:
    table_metadata_labels_by_section = _collect_table_metadata_labels_by_section(
        grid,
        sections,
    )
    return [
        section
        for section_index, section in enumerate(sections)
        if not (
            isinstance(section, dict)
            and _is_unanchored_synthetic_metadata_section(
                grid,
                section,
                section_index=section_index,
                table_metadata_labels_by_section=table_metadata_labels_by_section,
            )
        )
    ]


def _is_unanchored_synthetic_metadata_section(
    grid: Grid,
    section: dict[str, Any],
    *,
    section_index: int,
    table_metadata_labels_by_section: dict[int, set[str]],
) -> bool:
    if section.get("type") != "key_value":
        return False

    title = _optional_string(section.get("title"))
    if title is None or not _looks_like_synthetic_section_title(title):
        return False

    fields = section.get("fields")
    if not isinstance(fields, list) or not fields:
        return False

    field_labels = {_normalize_label(field).replace(" :", ":") for field in fields}
    if "" in field_labels or not all(
        isinstance(field, str) and _is_table_metadata_label(field)
        for field in fields
    ):
        return False

    if _find_section_title_cell(grid, title, start_row=1, end_row=len(grid)) is not None:
        return False

    previous_table_labels = _nearest_previous_table_metadata_labels(
        section_index,
        table_metadata_labels_by_section,
    )
    return previous_table_labels is not None and field_labels <= previous_table_labels


def _looks_like_synthetic_section_title(title: str) -> bool:
    return _SYNTHETIC_SECTION_TITLE_PATTERN.fullmatch(title.strip()) is not None


def _collect_table_metadata_labels_by_section(
    grid: Grid,
    sections: list[Any],
) -> dict[int, set[str]]:
    section_title_rows = _collect_section_title_rows(grid, sections)
    labels_by_section: dict[int, set[str]] = {}

    for section_index, section in enumerate(sections):
        if not isinstance(section, dict) or section.get("type") != "table":
            continue

        try:
            column_specs = _table_column_specs(section.get("headers", []))
            if not column_specs:
                continue

            start_row, end_row = _section_search_bounds(
                section_index,
                section_title_rows,
                row_count=len(grid),
            )
            header_row = _find_table_header_row(
                grid,
                column_specs,
                start_row=start_row,
                end_row=end_row,
            )
            if header_row is None:
                continue

            columns = []
            for column_spec in column_specs:
                column_index = header_row["columns_by_name"].get(
                    column_spec["match_key"]
                )
                if column_index is None:
                    continue

                columns.append(
                    {
                        "name": column_spec["output_name"],
                        "source_header": column_spec["source_name"],
                        "column_index": column_index,
                        "coordinate": _coordinate(header_row["row_index"], column_index),
                    }
                )

            if not columns:
                continue

            labels = _collect_table_metadata_labels(
                grid,
                header_row_index=header_row["row_index"],
                data_end_row_index=end_row,
                columns=columns,
            )
            if labels:
                labels_by_section[section_index] = labels
        except ExtractLocatorError:
            continue

    return labels_by_section


def _collect_table_metadata_labels(
    grid: Grid,
    *,
    header_row_index: int,
    data_end_row_index: int,
    columns: list[dict[str, Any]],
) -> set[str]:
    labels: set[str] = set()
    row_index = header_row_index + 1
    row_limit = min(data_end_row_index, len(grid))

    while row_index <= row_limit:
        row = grid[row_index - 1]
        metadata_row = _table_metadata_row(row, columns)
        if metadata_row is not None:
            labels.add(_normalize_label(metadata_row["name"]).replace(" :", ":"))

        row_index += 1

    return labels


def _nearest_previous_table_metadata_labels(
    section_index: int,
    table_metadata_labels_by_section: dict[int, set[str]],
) -> set[str] | None:
    for table_section_index in sorted(table_metadata_labels_by_section, reverse=True):
        if table_section_index < section_index:
            return table_metadata_labels_by_section[table_section_index]

    return None


def _locate_table_section(
    grid: Grid,
    section: dict[str, Any],
    *,
    section_index: int,
    section_title_rows: dict[int, int],
) -> dict[str, Any]:
    title = _optional_string(section.get("title"))
    column_specs = _table_column_specs(section.get("headers", []))
    if not column_specs:
        section_name = title or section_index + 1
        raise ExtractLocatorError(f'Table section "{section_name}" has no headers.')

    start_row, end_row = _section_search_bounds(
        section_index,
        section_title_rows,
        row_count=len(grid),
    )
    header_row = _find_table_header_row(
        grid,
        column_specs,
        start_row=start_row,
        end_row=end_row,
    )
    if header_row is None:
        raise ExtractLocatorError(
            f'Could not locate table header row for section "{title or section_index + 1}".'
        )

    located_columns = []
    for column_spec in column_specs:
        column_index = header_row["columns_by_name"].get(column_spec["match_key"])
        if column_index is None:
            continue

        located_columns.append(
            {
                "name": column_spec["output_name"],
                "source_header": column_spec["source_name"],
                "column_index": column_index,
                "coordinate": _coordinate(header_row["row_index"], column_index),
            }
        )

    if not located_columns:
        raise ExtractLocatorError(
            f'Could not locate any columns for table section "{title or section_index + 1}".'
        )

    return {
        "type": "table",
        "title": title,
        "output_title": _normalize_section_title_for_output(title),
        "section_index": section_index,
        "header_row_index": header_row["row_index"],
        "data_end_row_index": end_row,
        "columns": located_columns,
    }


def _locate_matrix_section(
    grid: Grid,
    section: dict[str, Any],
    *,
    section_index: int,
    section_title_rows: dict[int, int],
    merged_ranges: list[dict[str, int]],
) -> dict[str, Any]:
    title = _optional_string(section.get("title"))
    header_specs = _table_column_specs(section.get("headers", []))
    if not header_specs:
        section_name = title or section_index + 1
        raise ExtractLocatorError(f'Matrix section "{section_name}" has no headers.')

    start_row, end_row = _section_search_bounds(
        section_index,
        section_title_rows,
        row_count=len(grid),
    )
    candidates = detect_matrices(
        grid,
        start_row=start_row,
        end_row=end_row,
    )
    candidate = _find_matrix_candidate(
        candidates,
        title=title,
        expected_headers={spec["match_key"] for spec in header_specs},
    )
    if candidate is None:
        raise ExtractLocatorError(
            f'Could not locate matrix section "{title or section_index + 1}".'
        )

    headers_by_name = {
        _normalize_label(header["name"]): header
        for header in candidate["headers"]
    }
    located_columns: list[dict[str, Any]] = []
    for header_spec in header_specs:
        detected_header = headers_by_name.get(header_spec["match_key"])
        if detected_header is None:
            continue

        located_columns.append(
            {
                "name": header_spec["output_name"],
                "source_header": header_spec["source_name"],
                "start_column_index": detected_header["start_column_index"],
                "end_column_index": detected_header["end_column_index"],
                "coordinate": detected_header["coordinate"],
            }
        )

    if not located_columns:
        raise ExtractLocatorError(
            f'Could not locate any columns for matrix section "{title or section_index + 1}".'
        )

    row_header = _optional_string(section.get("row_header"))
    return {
        "type": "matrix",
        "title": title,
        "output_title": _normalize_section_title_for_output(title),
        "section_index": section_index,
        "header_row_index": candidate["header_row_index"],
        "row_label_column_index": candidate["row_label_column_index"],
        "data_start_row_index": candidate["data_start_row_index"],
        "data_end_row_index": candidate["data_end_row_index"],
        "columns": located_columns,
        "merged_ranges": _matrix_merged_ranges(
            merged_ranges,
            candidate=candidate,
            columns=located_columns,
        ),
        **({"row_header": row_header} if row_header is not None else {}),
    }


def _find_matrix_candidate(
    candidates: list[DetectedMatrix],
    *,
    title: str | None,
    expected_headers: set[str],
) -> DetectedMatrix | None:
    normalized_title = _normalize_label(title)
    if normalized_title:
        for candidate in candidates:
            if _normalize_label(candidate.get("title")) == normalized_title:
                return candidate

    best_candidate: DetectedMatrix | None = None
    best_overlap = 0
    for candidate in candidates:
        candidate_headers = {
            _normalize_label(header["name"])
            for header in candidate["headers"]
        }
        overlap = len(candidate_headers & expected_headers)
        required_overlap = min(2, len(expected_headers))
        if overlap >= required_overlap and overlap > best_overlap:
            best_candidate = candidate
            best_overlap = overlap

    return best_candidate


def _matrix_merged_ranges(
    merged_ranges: list[dict[str, int]],
    *,
    candidate: DetectedMatrix,
    columns: list[dict[str, Any]],
) -> list[dict[str, int]]:
    matrix_min_column = min(
        _require_int(column.get("start_column_index"), "matrix column start")
        for column in columns
    )
    matrix_max_column = max(
        _require_int(column.get("end_column_index"), "matrix column end")
        for column in columns
    )
    data_start_row = candidate["data_start_row_index"]
    data_end_row = candidate["data_end_row_index"]

    return [
        dict(merged_range)
        for merged_range in merged_ranges
        if merged_range["min_row"] <= data_end_row
        and merged_range["max_row"] >= data_start_row
        and merged_range["min_col"] <= matrix_max_column
        and merged_range["max_col"] >= matrix_min_column
    ]


def _collect_section_title_rows(
    grid: Grid,
    sections: list[Any],
) -> dict[int, int]:
    title_rows: dict[int, int] = {}
    for index, section in enumerate(sections):
        if not isinstance(section, dict):
            continue

        title = _optional_string(section.get("title"))
        if not title:
            continue

        title_cell = _find_section_title_cell(grid, title, start_row=1, end_row=len(grid))
        if title_cell is not None:
            title_rows[index] = title_cell["row_index"]

    return title_rows


def _section_search_bounds(
    section_index: int,
    section_title_rows: dict[int, int],
    *,
    row_count: int,
) -> tuple[int, int]:
    title_row = section_title_rows.get(section_index)
    start_row = title_row if title_row is not None else 1

    later_title_rows = [
        row_index
        for later_index, row_index in section_title_rows.items()
        if later_index > section_index and row_index > start_row
    ]
    if later_title_rows:
        end_row = min(later_title_rows) - 1
    else:
        end_row = min(row_count, start_row + _MAX_SECTION_SEARCH_ROWS)

    return max(1, start_row), max(start_row, end_row)


def _find_table_header_row(
    grid: Grid,
    column_specs: list[dict[str, Any]],
    *,
    start_row: int,
    end_row: int,
) -> dict[str, Any] | None:
    required_matches = 1 if len(column_specs) == 1 else max(2, len(column_specs) // 2)
    expected_keys = {column_spec["match_key"] for column_spec in column_specs}

    best_row: dict[str, Any] | None = None
    best_score = 0
    for row_index in range(start_row, min(end_row, len(grid)) + 1):
        row = grid[row_index - 1]
        columns_by_name = _first_matching_columns(row, expected_keys)
        score = len(columns_by_name)
        if score < required_matches:
            continue

        if score > best_score:
            best_row = {"row_index": row_index, "columns_by_name": columns_by_name}
            best_score = score

        if score == len(expected_keys):
            break

    return best_row


def _first_matching_columns(
    row: list[Any | None],
    expected_keys: set[str],
) -> dict[str, int]:
    columns_by_name: dict[str, int] = {}
    previous_normalized_value: str | None = None

    for column_offset, value in enumerate(row):
        normalized = _normalize_label(value)
        if not normalized:
            continue

        column_index = column_offset + 1
        if normalized == previous_normalized_value:
            continue

        previous_normalized_value = normalized
        if normalized in expected_keys and normalized not in columns_by_name:
            columns_by_name[normalized] = column_index

    return columns_by_name


def _table_column_specs(headers: Any) -> list[dict[str, str]]:
    if not isinstance(headers, list):
        return []

    specs: list[dict[str, str]] = []
    for header in headers:
        if not isinstance(header, dict):
            continue

        name = _optional_string(header.get("name"))
        if not name:
            continue

        subheaders = header.get("subheaders")
        if isinstance(subheaders, list) and subheaders:
            for subheader in subheaders:
                if not isinstance(subheader, dict):
                    continue

                subheader_name = _optional_string(subheader.get("name"))
                if subheader_name:
                    specs.append(_column_spec(subheader_name, source_name=name))
            continue

        specs.append(_column_spec(name, source_name=name))

    return specs


def _column_spec(output_name: str, *, source_name: str) -> dict[str, str]:
    return {
        "output_name": output_name,
        "source_name": source_name,
        "match_key": _normalize_label(output_name),
    }


def _extract_key_value_section(
    grid: Grid,
    section_locator: dict[str, Any],
) -> dict[str, Any]:
    values: dict[str, Any] = {}
    for field_locator in section_locator.get("fields", []):
        field_name = _require_string(field_locator.get("name"), "field locator name")
        if field_locator.get("status") != "found":
            values[field_name] = None
            continue

        value = _cell_at(
            grid,
            field_locator.get("value_row_index"),
            field_locator.get("value_column_index"),
        )
        values[field_name] = _json_safe_value(value)

    return values


def _extract_rebo_comment_section(
    grid: Grid,
    section_locator: dict[str, Any],
) -> dict[str, Any]:
    title_cell = _rebo_comment_title_cell(grid, section_locator)
    if title_cell is None:
        return {}

    comments: dict[str, Any] = {}
    title_row_index = title_cell["row_index"]
    title_column_index = title_cell["column_index"]

    for row_index in range(title_row_index + 1, len(grid) + 1):
        row = grid[row_index - 1]
        comment = _rebo_comment_value_for_row(row, title_column_index)
        if comment is None:
            continue

        comments[f"row_{row_index}"] = comment

    return comments


def _rebo_comment_title_cell(
    grid: Grid,
    section_locator: dict[str, Any],
) -> dict[str, Any] | None:
    row_index = section_locator.get("title_row_index")
    column_index = section_locator.get("title_column_index")
    if isinstance(row_index, int) and isinstance(column_index, int):
        return {
            "row_index": row_index,
            "column_index": column_index,
            "coordinate": _coordinate(row_index, column_index),
        }

    title = _optional_string(section_locator.get("title")) or _REBO_COMMENT_OUTPUT_TITLE
    return _find_section_title_cell(grid, title, start_row=1, end_row=len(grid))


def _rebo_comment_value_for_row(
    row: list[Any | None],
    title_column_index: int,
) -> Any | None:
    values: list[Any] = []
    seen: set[str] = set()

    for column_offset in range(max(title_column_index - 1, 0), len(row)):
        value = row[column_offset]
        if _is_empty(value):
            continue

        normalized = _normalize_label(value)
        if normalized in _REBO_COMMENT_NORMALIZED_TITLES:
            continue

        dedupe_key = normalized or str(value).strip()
        if dedupe_key in seen:
            continue

        seen.add(dedupe_key)
        values.append(_json_safe_value(value))

    if not values:
        return None

    if len(values) == 1:
        return values[0]

    return " | ".join(str(value) for value in values if value is not None)


def _extract_table_section(
    grid: Grid,
    section_locator: dict[str, Any],
) -> list[dict[str, Any]] | dict[str, Any]:
    header_row_index = _require_int(
        section_locator.get("header_row_index"),
        "table locator header_row_index",
    )
    columns = section_locator.get("columns")
    if not isinstance(columns, list) or not columns:
        raise ExtractLocatorError("Table locator must contain at least one column.")

    rows: list[dict[str, Any]] = []
    metadata: dict[str, Any] = {}
    row_index = header_row_index + 1
    data_end_row_index = section_locator.get("data_end_row_index")
    if not isinstance(data_end_row_index, int):
        data_end_row_index = len(grid)

    row_limit = min(data_end_row_index, len(grid))
    while row_index <= row_limit:
        row = grid[row_index - 1]

        if _is_blank_for_columns(row, columns):
            if rows:
                break
            row_index += 1
            continue

        if _looks_like_repeated_header(row, columns):
            row_index += 1
            continue

        metadata_row = _table_metadata_row(row, columns)
        if metadata_row is not None:
            metadata_key = _unique_key(metadata, metadata_row["name"])
            metadata[metadata_key] = _json_safe_value(metadata_row["value"])
            row_index += 1
            continue

        row_object: dict[str, Any] = {}
        for column in columns:
            name = _require_string(column.get("name"), "table column name")
            column_index = _require_int(column.get("column_index"), "table column index")
            row_object[name] = _json_safe_value(_cell_at(grid, row_index, column_index))

        if any(value is not None for value in row_object.values()):
            rows.append(row_object)

        row_index += 1

    if not metadata:
        return rows

    table_output: dict[str, Any] = {"rows": rows}
    for key, value in metadata.items():
        table_output[_unique_key(table_output, key)] = value

    return table_output


def _extract_matrix_section(
    grid: Grid,
    section_locator: dict[str, Any],
) -> dict[str, dict[str, Any]]:
    row_label_column_index = _require_int(
        section_locator.get("row_label_column_index"),
        "matrix locator row_label_column_index",
    )
    data_start_row_index = _require_int(
        section_locator.get("data_start_row_index"),
        "matrix locator data_start_row_index",
    )
    data_end_row_index = _require_int(
        section_locator.get("data_end_row_index"),
        "matrix locator data_end_row_index",
    )
    columns = section_locator.get("columns")
    if not isinstance(columns, list) or not columns:
        raise ExtractLocatorError("Matrix locator must contain at least one column.")
    merged_ranges = section_locator.get("merged_ranges", [])
    if not isinstance(merged_ranges, list):
        merged_ranges = []

    output: dict[str, dict[str, Any]] = {}
    row_limit = min(data_end_row_index, len(grid))
    for row_index in range(data_start_row_index, row_limit + 1):
        raw_label = _cell_at(grid, row_index, row_label_column_index)
        safe_label = _json_safe_value(raw_label)
        if safe_label is None:
            continue

        label = str(safe_label).strip()
        if not label:
            continue

        row_key = _unique_matrix_row_key(output, label)
        row_output: dict[str, Any] = {}
        for column in columns:
            name = _require_string(column.get("name"), "matrix column name")
            start_column_index = _require_int(
                column.get("start_column_index"),
                "matrix column start_column_index",
            )
            end_column_index = _require_int(
                column.get("end_column_index"),
                "matrix column end_column_index",
            )
            row_output[name] = _matrix_range_value(
                grid[row_index - 1],
                row_index=row_index,
                start_column_index=start_column_index,
                end_column_index=end_column_index,
                merged_ranges=merged_ranges,
            )

        output[row_key] = row_output

    return output


def _matrix_range_value(
    row: list[Any | None],
    *,
    row_index: int,
    start_column_index: int,
    end_column_index: int,
    merged_ranges: list[dict[str, Any]],
) -> Any:
    values: list[Any] = []
    seen: set[str] = set()

    for column_index in range(start_column_index, end_column_index + 1):
        if _is_non_anchor_merged_cell(
            row_index=row_index,
            column_index=column_index,
            merged_ranges=merged_ranges,
        ):
            continue

        safe_value = _json_safe_value(_cell_at_row(row, column_index))
        if safe_value is None:
            continue

        dedupe_key = _normalize_label(safe_value) or repr(safe_value)
        if dedupe_key in seen:
            continue

        seen.add(dedupe_key)
        values.append(safe_value)

    if not values:
        return None
    if len(values) == 1:
        return values[0]

    return " | ".join(str(value) for value in values)


def _is_non_anchor_merged_cell(
    *,
    row_index: int,
    column_index: int,
    merged_ranges: list[dict[str, Any]],
) -> bool:
    for merged_range in merged_ranges:
        try:
            min_row = _require_int(merged_range.get("min_row"), "merged min_row")
            max_row = _require_int(merged_range.get("max_row"), "merged max_row")
            min_col = _require_int(merged_range.get("min_col"), "merged min_col")
            max_col = _require_int(merged_range.get("max_col"), "merged max_col")
        except ExtractLocatorError:
            continue

        if (
            min_row <= row_index <= max_row
            and min_col <= column_index <= max_col
        ):
            return row_index != min_row

    return False


def _unique_matrix_row_key(
    existing: dict[str, Any],
    preferred_key: str,
) -> str:
    if preferred_key not in existing:
        return preferred_key

    index = 2
    while f"{preferred_key} {index}" in existing:
        index += 1

    return f"{preferred_key} {index}"


def _find_label_cell(
    grid: Grid,
    label: str,
    *,
    start_row: int,
    end_row: int,
) -> dict[str, Any] | None:
    normalized_label = _normalize_label(label)
    if not normalized_label:
        return None

    for row_index in range(start_row, min(end_row, len(grid)) + 1):
        row = grid[row_index - 1]
        previous_normalized_value: str | None = None
        for column_offset, value in enumerate(row):
            normalized = _normalize_label(value)
            if not normalized:
                continue

            if normalized == previous_normalized_value:
                continue

            previous_normalized_value = normalized
            if normalized == normalized_label:
                column_index = column_offset + 1
                return {
                    "row_index": row_index,
                    "column_index": column_index,
                    "coordinate": _coordinate(row_index, column_index),
                }

    return None


def _find_section_title_cell(
    grid: Grid,
    title: str,
    *,
    start_row: int,
    end_row: int,
) -> dict[str, Any] | None:
    exact_cell = _find_label_cell(grid, title, start_row=start_row, end_row=end_row)
    if exact_cell is not None:
        return exact_cell

    prefix = _section_title_match_prefix(title)
    if prefix is None:
        return None

    for row_index in range(start_row, min(end_row, len(grid)) + 1):
        row = grid[row_index - 1]
        previous_normalized_value: str | None = None
        for column_offset, value in enumerate(row):
            normalized = _normalize_label(value)
            if not normalized:
                continue

            if normalized == previous_normalized_value:
                continue

            previous_normalized_value = normalized
            if normalized == prefix or normalized.startswith(f"{prefix}:"):
                column_index = column_offset + 1
                return {
                    "row_index": row_index,
                    "column_index": column_index,
                    "coordinate": _coordinate(row_index, column_index),
                }

    return None


def _find_value_cell_to_right(
    grid: Grid,
    *,
    row_index: int,
    column_index: int,
    label: str,
    stop_labels: set[str],
) -> dict[str, Any] | None:
    row = grid[row_index - 1]
    normalized_label = _normalize_label(label)

    for column_offset in range(column_index, len(row)):
        value = row[column_offset]
        if _is_empty(value):
            continue

        normalized_value = _normalize_label(value)
        if normalized_value == normalized_label:
            continue

        if normalized_value in stop_labels:
            return None

        value_column_index = column_offset + 1
        return {
            "row_index": row_index,
            "column_index": value_column_index,
            "coordinate": _coordinate(row_index, value_column_index),
        }

    return None


def _looks_like_repeated_header(
    row: list[Any | None],
    columns: list[dict[str, Any]],
) -> bool:
    matched = 0
    for column in columns:
        name = _require_string(column.get("name"), "table column name")
        column_index = _require_int(column.get("column_index"), "table column index")
        if _normalize_label(_cell_at_row(row, column_index)) == _normalize_label(name):
            matched += 1

    return matched == len(columns)


def _table_metadata_row(
    row: list[Any | None],
    columns: list[dict[str, Any]],
) -> dict[str, Any] | None:
    ordered_columns = sorted(
        columns,
        key=lambda column: _require_int(
            column.get("column_index"),
            "table column index",
        ),
    )
    if not ordered_columns:
        return None

    first_column_index = _require_int(
        ordered_columns[0].get("column_index"),
        "table column index",
    )
    label = _cell_at_row(row, first_column_index)
    if not _is_table_metadata_label(label):
        return None

    label_text = _require_string(label, "table metadata label")
    normalized_label = _normalize_label(label_text)
    value: Any | None = None
    normalized_values: set[str] = set()

    for column in ordered_columns[1:]:
        column_index = _require_int(column.get("column_index"), "table column index")
        candidate = _cell_at_row(row, column_index)
        if _is_empty(candidate):
            continue

        normalized_candidate = _normalize_label(candidate)
        if normalized_candidate == normalized_label:
            continue

        if value is None:
            value = candidate

        normalized_values.add(normalized_candidate or str(candidate).strip().casefold())
        if len(normalized_values) > 1:
            return None

    return {"name": label_text, "value": value}


def _is_table_metadata_label(value: Any) -> bool:
    if not isinstance(value, str):
        return False

    label = _WHITESPACE_PATTERN.sub(" ", value).strip()
    return label.endswith(":")


def _is_blank_for_columns(
    row: list[Any | None],
    columns: list[dict[str, Any]],
) -> bool:
    for column in columns:
        column_index = _require_int(column.get("column_index"), "table column index")
        if not _is_empty(_cell_at_row(row, column_index)):
            return False

    return True


def _find_sheet_grid(
    sheets: dict[str, Grid],
    requested_sheet_name: str,
) -> tuple[str, Grid]:
    if requested_sheet_name in sheets:
        return requested_sheet_name, sheets[requested_sheet_name]

    normalized_requested = _normalize_label(requested_sheet_name)
    for sheet_name, grid in sheets.items():
        if _normalize_label(sheet_name) == normalized_requested:
            return sheet_name, grid

    raise ExtractLocatorError(f'Sheet not found in workbook: "{requested_sheet_name}"')


def _section_output_title(section_locator: dict[str, Any]) -> str:
    output_title = _optional_string(section_locator.get("output_title"))
    if output_title:
        return output_title

    title = _optional_string(section_locator.get("title"))
    if title:
        return _normalize_section_title_for_output(title)

    section_index = section_locator.get("section_index")
    if isinstance(section_index, int):
        return f"Section {section_index + 1}"

    return "Section"


def _is_rebo_comment_section(section_locator: dict[str, Any]) -> bool:
    output_title = _optional_string(section_locator.get("output_title"))
    if output_title and _normalize_label(output_title) in _REBO_COMMENT_NORMALIZED_TITLES:
        return True

    title = _optional_string(section_locator.get("title"))
    return bool(title and _normalize_label(title) in _REBO_COMMENT_NORMALIZED_TITLES)


def _unique_key(existing: dict[str, Any], preferred_key: str) -> str:
    if preferred_key not in existing:
        return preferred_key

    index = 2
    while f"{preferred_key} ({index})" in existing:
        index += 1

    return f"{preferred_key} ({index})"


def _is_supported_field_locators(field_locators: Any) -> bool:
    return (
        isinstance(field_locators, dict)
        and field_locators.get("version") == _LOCATOR_VERSION
        and field_locators.get("type") == _LOCATOR_TYPE
        and isinstance(field_locators.get("sheets"), list)
    )


def _is_supported_document_field_locators(
    field_locators: Any,
    schema_id: str,
) -> bool:
    return (
        _is_supported_field_locators(field_locators)
        and field_locators.get("schema_id") == schema_id
    )


def _section_title_match_prefix(title: str) -> str | None:
    normalized_title = _normalize_label(title)
    if normalized_title.startswith("lease information:"):
        return "lease information"

    return None


def _normalize_section_title_for_output(title: str | None) -> str | None:
    if title is None:
        return None

    if _section_title_match_prefix(title) == "lease information":
        return "Lease Information"

    return title


def _key_value_section_output_title(
    grid: Grid,
    title: str | None,
    field_locators: list[dict[str, Any]],
) -> str | None:
    if title and _looks_like_synthetic_section_title(title):
        inferred_title = _infer_synthetic_key_value_section_title(
            grid,
            field_locators,
        )
        if inferred_title is not None:
            return inferred_title

    return _normalize_section_title_for_output(title)


def _infer_synthetic_key_value_section_title(
    grid: Grid,
    field_locators: list[dict[str, Any]],
) -> str | None:
    candidate_counts: dict[str, int] = {}
    candidate_texts: dict[str, str] = {}

    for field_locator in field_locators:
        row_index = field_locator.get("label_row_index")
        label_column_index = field_locator.get("label_column_index")
        if not isinstance(row_index, int) or not isinstance(label_column_index, int):
            continue

        row = grid[row_index - 1]
        seen_on_row: set[str] = set()
        for column_index in range(1, label_column_index):
            value = _cell_at_row(row, column_index)
            normalized = _normalize_label(value)
            if not normalized or normalized in seen_on_row:
                continue

            seen_on_row.add(normalized)
            safe_value = _json_safe_value(value)
            if safe_value is None:
                continue

            text = str(safe_value).strip()
            if not text or _looks_like_synthetic_section_title(text):
                continue

            candidate_counts[normalized] = candidate_counts.get(normalized, 0) + 1
            candidate_texts.setdefault(normalized, text)

    if not candidate_counts:
        return None

    best_normalized = max(
        candidate_counts,
        key=lambda normalized: (
            candidate_counts[normalized],
            len(candidate_texts[normalized]),
        ),
    )
    return _normalize_section_title_for_output(candidate_texts[best_normalized])


def _resolve_stored_path(stored_path: str) -> Path:
    relative_path = Path(stored_path)
    if relative_path.is_absolute():
        raise ExtractStoredDocumentFileError("Stored document path must be relative.")

    resolved_path = (_RESOLVED_BACKEND_ROOT / relative_path).resolve()
    try:
        resolved_path.relative_to(_RESOLVED_BACKEND_ROOT)
    except ValueError as error:
        raise ExtractStoredDocumentFileError("Stored document path escapes backend root.") from error

    if not resolved_path.is_file():
        raise ExtractStoredDocumentFileError(f"Stored document file not found: {stored_path}")

    return resolved_path


def _load_workbook_sheets(path: Path) -> tuple[dict[str, Grid], MergedRanges]:
    try:
        return load_sheets_with_merged_ranges(path)
    except FileNotFoundError as error:
        raise ExtractStoredDocumentFileError(str(error)) from error
    except ValueError as error:
        raise ExtractInvalidDocumentFileError(str(error)) from error


def _cell_at(grid: Grid, row_index: Any, column_index: Any) -> Any | None:
    if not isinstance(row_index, int) or not isinstance(column_index, int):
        return None

    if row_index < 1 or row_index > len(grid):
        return None

    return _cell_at_row(grid[row_index - 1], column_index)


def _cell_at_row(row: list[Any | None], column_index: int) -> Any | None:
    if column_index < 1:
        return None

    column_offset = column_index - 1
    if column_offset >= len(row):
        return None

    return row[column_offset]


def _is_placeholder_only_value(value: Any) -> bool:
    if not isinstance(value, str):
        return False

    if not _PLACEHOLDER_TOKEN_PATTERN.search(value):
        return False

    remainder = _PLACEHOLDER_TOKEN_PATTERN.sub("", value)
    return _ALNUM_PATTERN.search(remainder) is None


def _json_safe_value(value: Any) -> Any:
    if value is None:
        return None

    if _is_placeholder_only_value(value):
        return None

    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == time.min else value.isoformat()

    if isinstance(value, date):
        return value.isoformat()

    if isinstance(value, time):
        return value.isoformat()

    if isinstance(value, Decimal):
        return float(value)

    return value


def _is_empty(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _normalize_label(value: Any) -> str:
    if not isinstance(value, str):
        return ""

    return _WHITESPACE_PATTERN.sub(" ", value).strip().casefold()


def _optional_string(value: Any) -> str | None:
    if not isinstance(value, str):
        return None

    stripped = value.strip()
    return stripped or None


def _require_string(value: Any, name: str) -> str:
    stripped = _optional_string(value)
    if stripped is None:
        raise ExtractLocatorError(f"{name} must be a non-empty string.")

    return stripped


def _require_int(value: Any, name: str) -> int:
    if not isinstance(value, int):
        raise ExtractLocatorError(f"{name} must be an integer.")

    return value


def _coordinate(row_index: int, column_index: int) -> str:
    return f"{get_column_letter(column_index)}{row_index}"
