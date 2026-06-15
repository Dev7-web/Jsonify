from typing import Any,Dict,List,Tuple
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from io import BytesIO
import re

from app.schemas.excel_schema import DetectedHeader, HeaderDetectionResult

ExcelCellValue=Any

def normalize_cell_value(value: ExcelCellValue) -> str:
    if value is None:
        return ""
    
    return re.sub(r"\s+"," ", str(value).strip())


def normalize_header_name(value: str) -> str:
    return re.sub(r"\s+"," ",value.strip())

def has_alphabet(value:str) -> bool:
    return bool(re.search(r"[a-zA-Z]",value))

def is_number_like(value:str) -> bool:
    if not value:
        return False
    
    try:
        float(value)
        return True
    except ValueError:
        return False
    

def is_date_like(value: str) -> bool:
    if not value:
        return False

    date_patterns = [
        r"^\d{1,2}[-/]\d{1,2}[-/]\d{2,4}$",
        r"^\d{4}[-/]\d{1,2}[-/]\d{1,2}$",
    ]

    return any(re.match(pattern, value) for pattern in date_patterns)


def get_non_empty_cells(row: List[ExcelCellValue]) -> List[str]:
    return [
        normalize_cell_value(cell)
        for cell in row
        if normalize_cell_value(cell)
    ]
    
def calculate_row_score(
    row: List[ExcelCellValue],
    all_rows: List[List[ExcelCellValue]],
    row_index: int
) -> int:
    
    cells = [normalize_cell_value(cell) for cell in row]
    non_empty_cells = [cell for cell in cells if cell]

    if len(non_empty_cells) < 2:
        return 0

    score = 0

    text_cells = [cell for cell in non_empty_cells if has_alphabet(cell)]
    number_cells = [cell for cell in non_empty_cells if is_number_like(cell)]
    date_cells = [cell for cell in non_empty_cells if is_date_like(cell)]

    unique_cells = set(cell.lower() for cell in non_empty_cells)

    unique_ratio = len(unique_cells) / len(non_empty_cells)
    text_ratio = len(text_cells) / len(non_empty_cells)

    # Rule 1:
    # Header row usually has more text.
    score += len(text_cells) * 10

    # Rule 2:
    # Header row should not be mostly numbers.
    score -= len(number_cells) * 8

    # Rule 3:
    # Header row should not be mostly dates.
    score -= len(date_cells) * 8

    # Rule 4:
    # Header names are usually unique.
    if unique_ratio >= 0.8:
        score += 20
    else:
        score -= 20

    # Rule 5:
    # At least 60% cells should look like text.
    if text_ratio >= 0.6:
        score += 25
    else:
        score -= 10

    # Rule 6:
    # Very long text usually means title/description, not header.
    very_long_cells = [cell for cell in non_empty_cells if len(cell) > 40]
    score -= len(very_long_cells) * 10

    # Rule 7:
    # Data should exist below the header row.
    next_rows = all_rows[row_index + 1: row_index + 6]

    useful_rows_below = 0

    for next_row in next_rows:
        next_non_empty_cells = get_non_empty_cells(next_row)

        if len(next_non_empty_cells) >= max(2, len(non_empty_cells) * 0.5):
            useful_rows_below += 1

    score += useful_rows_below * 10

    return score

def find_best_header_row(
    rows: List[List[ExcelCellValue]]
) -> Tuple[int, int]:
    max_rows_to_check = min(len(rows), 30)

    best_row_index = -1
    best_score = 0

    for index in range(max_rows_to_check):
        score = calculate_row_score(rows[index], rows, index)

        if score > best_score:
            best_score = score
            best_row_index = index

    if best_row_index == -1:
        raise Exception("Could not detect header row")

    confidence_score = min(100, max(0, best_score))

    return best_row_index, confidence_score


def build_headers(header_row: List[ExcelCellValue]) -> List[DetectedHeader]:
    headers: List[DetectedHeader] = []

    for column_index, cell in enumerate(header_row):
        header_name = normalize_header_name(normalize_cell_value(cell))

        if not header_name:
            continue

        headers.append(
            DetectedHeader(
                name=header_name,
                column_index=column_index,
                excel_column_name=get_column_letter(column_index + 1)
            )
        )

    return headers

def build_preview_rows(
    rows: List[List[ExcelCellValue]],
    header_row_index: int,
    headers: List[DetectedHeader]
) -> List[Dict[str, ExcelCellValue]]:
    data_rows = rows[header_row_index + 1: header_row_index + 6]

    preview_rows: List[Dict[str, ExcelCellValue]] = []

    for row in data_rows:
        row_object: Dict[str, ExcelCellValue] = {}

        for header in headers:
            value = ""

            if header.column_index < len(row):
                value = row[header.column_index]

            row_object[header.name] = value

        preview_rows.append(row_object)

    return preview_rows


def detect_headers_from_excel(file_bytes: bytes) -> HeaderDetectionResult:
    workbook = load_workbook(
        filename=BytesIO(file_bytes),
        data_only=True
    )

    first_sheet_name = workbook.sheetnames[0]

    if not first_sheet_name:
        raise Exception("Excel file does not contain any sheet")

    worksheet = workbook[first_sheet_name]

    rows: List[List[ExcelCellValue]] = []

    for row in worksheet.iter_rows(values_only=True):
        row_values = list(row)

        if any(cell is not None and str(cell).strip() != "" for cell in row_values):
            rows.append(row_values)

    if len(rows) == 0:
        raise Exception("Excel sheet is empty")

    header_row_index, confidence_score = find_best_header_row(rows)

    header_row = rows[header_row_index]

    headers = build_headers(header_row)

    preview_rows = build_preview_rows(rows, header_row_index, headers)

    return HeaderDetectionResult(
        sheet_name=first_sheet_name,
        header_row_index=header_row_index,
        header_row_number=header_row_index + 1,
        confidence_score=confidence_score,
        headers=headers,
        preview_rows=preview_rows
    )
    
    

    
    
    
    
    
    

    
    
    
    


