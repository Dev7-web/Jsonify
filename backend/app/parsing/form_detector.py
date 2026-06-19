from __future__ import annotations

from pathlib import Path
from typing import Any, TypedDict

from openpyxl.worksheet.worksheet import Worksheet

from .excel_loader import _open_workbook, normalize_cell_value


class DetectedKeyValueField(TypedDict):
    name: str
    coordinate: str


class DetectedKeyValueSection(TypedDict):
    title: str
    title_coordinate: str
    row_index: int
    coordinate: str
    fields: list[DetectedKeyValueField]


class _CellBlock(TypedDict):
    value: Any
    text: str
    coordinate: str
    start_column: int
    end_column: int
    solid_fill: bool


_MAX_SECTION_SCAN_ROWS = 80
_MIN_FIELDS_TO_TRUST_SECTION = 2
_MIN_TRUSTED_SECTIONS_PER_BAND_ROW = 1


def detect_key_value_sections(
    path: str | Path,
) -> dict[str, list[DetectedKeyValueSection]]:
    """Detect form-style key/value sections from Excel title bands.

    This parser uses workbook structure that the LLM flattened text cannot
    reliably preserve: merged horizontal title bands and filled title cells.
    """
    workbook = _open_workbook(path, read_only=False)

    try:
        return {
            worksheet.title: _detect_sheet_key_value_sections(worksheet)
            for worksheet in workbook.worksheets
        }
    finally:
        workbook.close()


def _detect_sheet_key_value_sections(
    worksheet: Worksheet,
) -> list[DetectedKeyValueSection]:
    merged_spans = _merged_column_spans_by_top_left(worksheet)
    sections: list[DetectedKeyValueSection] = []

    for row_index in range(1, worksheet.max_row + 1):
        title_blocks = _title_blocks_for_row(worksheet, row_index, merged_spans)
        if len(title_blocks) < 2:
            continue

        stop_row = _find_section_scan_stop_row(worksheet, row_index, merged_spans)
        candidates = [
            _build_key_value_section(worksheet, title_block, row_index, stop_row, merged_spans)
            for title_block in title_blocks
        ]
        trusted_count = sum(
            1
            for candidate in candidates
            if len(candidate["fields"]) >= _MIN_FIELDS_TO_TRUST_SECTION
        )
        if trusted_count < _MIN_TRUSTED_SECTIONS_PER_BAND_ROW:
            continue

        sections.extend(candidates)

    return _deduplicate_sections(sections)


def _build_key_value_section(
    worksheet: Worksheet,
    title_block: _CellBlock,
    title_row_index: int,
    stop_row: int,
    merged_spans: dict[tuple[int, int], int],
) -> DetectedKeyValueSection:
    fields: list[DetectedKeyValueField] = []
    seen_fields: set[str] = set()

    for row_index in range(title_row_index + 1, stop_row + 1):
        field = _detect_field_in_row(
            worksheet,
            row_index,
            title_block["start_column"],
            title_block["end_column"],
            merged_spans,
        )
        if field is None:
            continue

        normalized_name = _normalize_label(field["name"])
        if normalized_name in seen_fields:
            continue

        seen_fields.add(normalized_name)
        fields.append(field)

    title = _clean_text(title_block["text"])
    return {
        "title": title,
        "title_coordinate": title_block["coordinate"],
        "row_index": title_row_index,
        "coordinate": title_block["coordinate"],
        "fields": fields,
    }


def _detect_field_in_row(
    worksheet: Worksheet,
    row_index: int,
    start_column: int,
    end_column: int,
    merged_spans: dict[tuple[int, int], int],
) -> DetectedKeyValueField | None:
    blocks = [
        block
        for block in _row_blocks(worksheet, row_index, merged_spans)
        if _overlaps(block, start_column, end_column)
    ]
    if len(blocks) < 2:
        return None

    label_block = blocks[0]
    if label_block["solid_fill"] or not isinstance(label_block["value"], str):
        return None

    if not any(block["start_column"] > label_block["end_column"] for block in blocks[1:]):
        return None

    label = _clean_text(label_block["text"])
    if not _is_probable_field_label(label):
        return None

    return {"name": label, "coordinate": label_block["coordinate"]}


def _title_blocks_for_row(
    worksheet: Worksheet,
    row_index: int,
    merged_spans: dict[tuple[int, int], int],
) -> list[_CellBlock]:
    return [
        block
        for block in _row_blocks(worksheet, row_index, merged_spans)
        if block["solid_fill"] and isinstance(block["value"], str) and _clean_text(block["text"])
    ]


def _find_section_scan_stop_row(
    worksheet: Worksheet,
    title_row_index: int,
    merged_spans: dict[tuple[int, int], int],
) -> int:
    max_scan_row = min(worksheet.max_row, title_row_index + _MAX_SECTION_SCAN_ROWS)

    for row_index in range(title_row_index + 1, max_scan_row + 1):
        if _title_blocks_for_row(worksheet, row_index, merged_spans):
            return row_index - 1

    return max_scan_row


def _row_blocks(
    worksheet: Worksheet,
    row_index: int,
    merged_spans: dict[tuple[int, int], int],
) -> list[_CellBlock]:
    blocks: list[_CellBlock] = []
    column_index = 1

    while column_index <= worksheet.max_column:
        cell = worksheet.cell(row_index, column_index)
        value = normalize_cell_value(cell.value)
        end_column = merged_spans.get((row_index, column_index), column_index)

        if value is not None:
            blocks.append(
                {
                    "value": value,
                    "text": str(value),
                    "coordinate": cell.coordinate,
                    "start_column": column_index,
                    "end_column": end_column,
                    "solid_fill": cell.fill.fill_type == "solid",
                }
            )

        column_index = end_column + 1

    return blocks


def _merged_column_spans_by_top_left(worksheet: Worksheet) -> dict[tuple[int, int], int]:
    return {
        (merged_range.min_row, merged_range.min_col): merged_range.max_col
        for merged_range in worksheet.merged_cells.ranges
    }


def _deduplicate_sections(
    sections: list[DetectedKeyValueSection],
) -> list[DetectedKeyValueSection]:
    deduplicated: list[DetectedKeyValueSection] = []
    seen: set[tuple[str, str]] = set()

    for section in sections:
        key = (_normalize_label(section["title"]), section["title_coordinate"])
        if key in seen:
            continue

        seen.add(key)
        deduplicated.append(section)

    return deduplicated


def _overlaps(block: _CellBlock, start_column: int, end_column: int) -> bool:
    return block["start_column"] <= end_column and block["end_column"] >= start_column


def _is_probable_field_label(value: str) -> bool:
    if not value:
        return False

    if len(value) > 80:
        return False

    return any(character.isalpha() for character in value)


def _clean_text(value: str) -> str:
    return " ".join(value.split()).strip()


def _normalize_label(value: str) -> str:
    return _clean_text(value).casefold()
