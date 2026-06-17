import re
import tempfile
from pathlib import Path
from typing import Any
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.worksheet import Worksheet


Grid = list[list[Any | None]]

_WHITESPACE_PATTERN = re.compile(r"\s+")
_UNSUPPORTED_STYLE_XXID_PATTERN = re.compile(rb"\sxxid=(\"[^\"]*\"|'[^']*')")
_STYLES_XML_PATH = "xl/styles.xml"


def normalize_cell_value(value: Any) -> Any | None:
    if value is None:
        return None

    if isinstance(value, str):
        normalized = _WHITESPACE_PATTERN.sub(" ", value).strip()
        return normalized or None

    return value


def _expand_merged_cells(worksheet: Worksheet) -> None:
    # Copy each merged range's top-left value into every cell in the range,
    # so downstream consumers (header detection) see merged headers like
    # `Q1 2024` spanning A1:D1 as four populated cells instead of one + Nones.
    for merged_range in list(worksheet.merged_cells.ranges):
        top_left_value = worksheet.cell(merged_range.min_row, merged_range.min_col).value
        worksheet.unmerge_cells(str(merged_range))
        for row_index in range(merged_range.min_row, merged_range.max_row + 1):
            for column_index in range(merged_range.min_col, merged_range.max_col + 1):
                worksheet.cell(row_index, column_index).value = top_left_value


def _open_workbook(path: str | Path, *, read_only: bool):
    workbook_path = Path(path)

    if not workbook_path.is_file():
        raise FileNotFoundError(f"Excel file not found: {workbook_path}")

    try:
        return load_workbook(workbook_path, data_only=True, read_only=read_only)
    except TypeError as error:
        if not _is_unsupported_xxid_style_error(error):
            raise

        return _open_workbook_with_repaired_styles(workbook_path, read_only=read_only)
    except (BadZipFile, InvalidFileException) as error:
        raise ValueError(f"File is not a valid .xlsx workbook: {workbook_path}") from error


def _is_unsupported_xxid_style_error(error: TypeError) -> bool:
    return "unexpected keyword argument 'xxid'" in str(error)


def _open_workbook_with_repaired_styles(path: Path, *, read_only: bool):
    repaired_path = _create_xxid_repaired_workbook(path)

    try:
        workbook = load_workbook(repaired_path, data_only=True, read_only=read_only)
    except (BadZipFile, InvalidFileException, TypeError) as error:
        repaired_path.unlink(missing_ok=True)
        raise ValueError(
            f"Excel workbook contains unsupported style metadata and could not be repaired: {path}"
        ) from error
    except Exception:
        repaired_path.unlink(missing_ok=True)
        raise

    return _cleanup_temp_workbook_on_close(workbook, repaired_path)


def _create_xxid_repaired_workbook(path: Path) -> Path:
    try:
        with ZipFile(path, "r") as source:
            if _STYLES_XML_PATH not in source.namelist():
                raise ValueError("Workbook does not contain xl/styles.xml.")

            with tempfile.NamedTemporaryFile(
                suffix=".xlsx",
                prefix="repaired-",
                delete=False,
            ) as temp_file:
                repaired_path = Path(temp_file.name)

            repaired = False
            try:
                with ZipFile(repaired_path, "w") as target:
                    for item in source.infolist():
                        data = source.read(item.filename)
                        if item.filename == _STYLES_XML_PATH:
                            data, replacements = _UNSUPPORTED_STYLE_XXID_PATTERN.subn(
                                b"",
                                data,
                            )
                            repaired = replacements > 0

                        target.writestr(item, data)
            except Exception:
                repaired_path.unlink(missing_ok=True)
                raise

            if not repaired:
                repaired_path.unlink(missing_ok=True)
                raise ValueError("Workbook styles did not contain unsupported xxid metadata.")

            return repaired_path
    except BadZipFile as error:
        raise ValueError(f"File is not a valid .xlsx workbook: {path}") from error


def _cleanup_temp_workbook_on_close(workbook, repaired_path: Path):
    original_close = workbook.close

    def close_with_cleanup() -> None:
        try:
            original_close()
        finally:
            repaired_path.unlink(missing_ok=True)

    workbook.close = close_with_cleanup
    return workbook


# ponytail: full grid kept in memory (no read_only mode, since merged-cell info
# is only available in normal mode). Safe under the 25 MB upload limit.
# Switch to streaming header detection if uploads grow much larger.
def load_sheets(path: str | Path) -> dict[str, Grid]:
    workbook = _open_workbook(path, read_only=False)

    try:
        sheets: dict[str, Grid] = {}

        for worksheet in workbook.worksheets:
            _expand_merged_cells(worksheet)

            grid: Grid = []

            for row in worksheet.iter_rows(values_only=True):
                grid.append([normalize_cell_value(cell) for cell in row])

            sheets[worksheet.title] = grid

        return sheets
    finally:
        workbook.close()
